#!/usr/bin/env python3
"""Generate a cover letter for the de novo resubmission."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

PROJECT_DIR = Path(__file__).resolve().parent.parent
OUT_DIR = PROJECT_DIR / "manuscript"
OUT_DIR.mkdir(parents=True, exist_ok=True)


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    p = doc.add_paragraph("To the Editors")
    p.paragraph_format.space_after = Pt(6)
    doc.add_paragraph("Public Choice").paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    p.add_run("Re: Network Exclusion and State Collapse: A Source-Anchored Reconstruction").bold = True
    p.paragraph_format.space_after = Pt(12)

    doc.add_paragraph("Dear Editors,").paragraph_format.space_after = Pt(12)

    body = (
        "We are pleased to submit the manuscript \"Network Exclusion and State Collapse: A Source-Anchored Reconstruction\" "
        "for consideration as a de novo submission. The paper has been substantially rebuilt in response to the concerns raised "
        "in the previous review process (Submission ID 32411b1b-3a61-4538-b0bd-c25fe1dc6692), and we believe it now provides "
        "a transparent, reproducible treatment of the relationship between network exclusion and state collapse.\n\n"
        "The key revisions are:\n"
        "- The full empirical pipeline is now source-anchored and reproducible from public data and code. "
        "All variables, figures, tables, and in-text numbers are generated from open repositories (Natural Earth, Nunn-Puga ruggedness, "
        "Our World in Data, Correlates of War, Cliopatria/Seshat) via version-controlled scripts, not hand-coded values.\n"
        "- A draft closure-type classification (maritime ban, land isolation, technology-network exclusion, bloc, patron-open, open, uncertain) "
        "is applied blind to outcome and documented with URLs or printed sources and confidence ratings.\n"
        "- Because the analysis sample is small (n=77), we report extensive small-sample robustness checks: "
        "Firth-penalized logistic regression, Fisher exact tests, exact-style permutation p-values, bootstrap confidence intervals, "
        "propensity-score weighted collapse rates, and a discrete-time survival expansion using a 50-year panel.\n"
        "- All reported numbers are rounded to three decimal places, all figures and tables are cited in the body in order of appearance, "
        "and references are numbered in Vancouver order and verified against Crossref/OpenLibrary.\n\n"
        "The manuscript remains a cross-sectional, observational study, and we accordingly frame the closure-type findings as "
        "associations rather than causal estimates. We would welcome the opportunity to revise further in light of your and the reviewers' comments."
    )
    for para in body.split("\n\n"):
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(12)

    doc.add_paragraph("Sincerely,").paragraph_format.space_after = Pt(24)
    doc.add_paragraph("[Author Name]")
    doc.add_paragraph("[Institutional affiliation]")
    doc.add_paragraph("[Email address]")

    out = OUT_DIR / "cover_letter.docx"
    doc.save(out)
    print(f"Cover letter written to {out}")


if __name__ == "__main__":
    main()
